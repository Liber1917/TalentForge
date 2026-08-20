"""resume_io 转接器测试：格式分派、PDF/DOCX 抽取、未知格式报错。"""

from __future__ import annotations

from pathlib import Path

import pytest

from talentforge.profile.resume_io import (
    UnsupportedResumeFormatError,
    extract_resume_text,
)


def test_txt_reads_directly(tmp_path: Path):
    p = tmp_path / "resume.txt"
    p.write_text("做过分布式存储课程项目", encoding="utf-8")
    assert extract_resume_text(p) == "做过分布式存储课程项目"


def test_md_reads_directly(tmp_path: Path):
    p = tmp_path / "resume.md"
    p.write_text("# 张三\n后端方向", encoding="utf-8")
    assert extract_resume_text(p) == "# 张三\n后端方向"


def test_pdf_extracts_text_layer(tmp_path: Path):
    from pypdf import PdfWriter

    # 空白页 PDF：抽取不抛异常且返回 str（空白页 → 空串是合法行为；
    # 含文本层的真实 PDF 抽取由真实简历冒烟覆盖）
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    p = tmp_path / "resume.pdf"
    with open(p, "wb") as f:
        writer.write(f)
    text = extract_resume_text(p)
    assert isinstance(text, str)


def test_docx_extracts_paragraphs(tmp_path: Path):
    from docx import Document

    p = tmp_path / "resume.docx"
    doc = Document()
    doc.add_paragraph("张三")
    doc.add_paragraph("负责后端服务研发")
    doc.add_paragraph("")  # 空段应被跳过
    doc.save(str(p))
    assert extract_resume_text(p) == "张三\n负责后端服务研发"


def test_unknown_suffix_raises(tmp_path: Path):
    p = tmp_path / "resume.xyz"
    p.write_text("x", encoding="utf-8")
    with pytest.raises(UnsupportedResumeFormatError):
        extract_resume_text(p)


def test_missing_file_raises(tmp_path: Path):
    with pytest.raises(FileNotFoundError):
        extract_resume_text(tmp_path / "nope.txt")
